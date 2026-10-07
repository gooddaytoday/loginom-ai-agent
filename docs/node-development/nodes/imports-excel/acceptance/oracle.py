"""Independent XLSX oracle. Never imports Loginom runtime or a node handler."""
import argparse
import datetime
import hashlib
import json
from pathlib import Path
import zipfile
import xml.etree.ElementTree as ET

import openpyxl

ROOT = Path(__file__).parent


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def expected_large(path):
    book = openpyxl.load_workbook(path, data_only=True, read_only=True)
    assert book.sheetnames == ['Data']
    rows = list(book['Data'].iter_rows(values_only=True))
    assert list(rows.pop(0)) == ['Id', 'Text', 'Amount', 'Occurred', 'Active']
    assert len(rows) == 20789
    result = []
    for i, row in enumerate(rows, 1):
        # This arithmetic specification is independent of both the writer and Loginom.
        wanted = (i, 'row-' + str(i) if i % 97 else None, i / 4,
                  datetime.datetime(2026, 1, 1) + datetime.timedelta(seconds=i), bool(i % 2))
        assert row == wanted, ('SOURCE_VALUE_MISMATCH', i)
        result.append([{'is_null': v is None, 'value': v.isoformat(timespec='milliseconds')
                        if isinstance(v, datetime.datetime) else v} for v in row])
    book.close()
    names = ['Id', 'Text', 'Amount', 'Occurred', 'Active']
    return {'row_count': 20789, 'schema': [dict(name=n, label=n, type=t)
            for n, t in zip(names, [4, 5, 3, 2, 1])], 'rows': result}


def verify_fixtures(directory):
    manifest = json.loads((directory / 'manifest.json').read_text())
    for item in manifest['files']:
        path = directory / item['name']
        assert path.stat().st_size == item['bytes'], ('SIZE_MISMATCH', item['name'])
        assert digest(path) == item['sha256'], ('SHA_MISMATCH', item['name'])
        assert item['bytes'] <= 16 * 1024 * 1024
    expected_large(directory / 'large.xlsx')
    ns = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
    with zipfile.ZipFile(directory / 'irregular.xlsx') as archive:
        xml = ET.fromstring(archive.read('xl/worksheets/sheet1.xml'))
        cells = {c.attrib['r']: c for c in xml.findall('.//s:c', ns)}
        assert cells['F2'].find('s:v', ns).text == '3'
        assert cells['F3'].find('s:f', ns).text == '1+2'
        assert cells['F3'].find('s:v', ns) is None
        assert cells['F4'].find('s:v', ns).text == '#DIV/0!'
        assert cells['F4'].attrib['t'] == 'e'
        assert xml.find('.//s:mergeCell', ns).attrib['ref'] == 'B5:C5'
        assert cells['B3'].attrib['t'] == 's'  # empty string is physically present
    return manifest


def compare(actual, expected, node_id):
    assert actual['node_id'] == node_id, 'WRONG_NODE'
    assert actual['port'] == 0, 'WRONG_PORT'
    assert actual['execution']['status'] == 3 and actual['execution']['error'] == ''
    assert actual['binding_rechecked'] is True
    assert actual['coverage'] == dict(rows=20789, cells=103945, complete=True, ordered=True)
    assert actual['row_count'] == expected['row_count']
    assert actual['schema'] == expected['schema'], 'SCHEMA_MISMATCH'
    assert len(actual['rows']) == len(expected['rows']), 'MISSING_ROWS'
    for i, (got, wanted) in enumerate(zip(actual['rows'], expected['rows']), 1):
        assert got == wanted, ('CELL_OR_ORDER_MISMATCH', i)
        for cell, expected_cell in zip(got, wanted):
            # Python otherwise considers False == 0; reject that conflation.
            assert type(cell['value']) is type(expected_cell['value']), ('VALUE_TYPE', i)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--fixtures', type=Path, default=ROOT / 'fixtures')
    parser.add_argument('--actual', type=Path, action='append', default=[])
    parser.add_argument('--node-id')
    parser.add_argument('--server-source', type=Path, action='append', default=[])
    args = parser.parse_args()
    manifest = verify_fixtures(args.fixtures)
    expected = expected_large(args.fixtures / 'large.xlsx')
    for path in args.server_source:
        assert digest(path) == digest(args.fixtures / 'large.xlsx'), 'SERVER_BYTES_MISMATCH'
    for path in args.actual:
        assert args.node_id, 'EXACT_NODE_REQUIRED'
        compare(json.loads(path.read_text()), expected, args.node_id)
    print(json.dumps({'status': 'PASS', 'fixtures': len(manifest['files']),
                      'audits': len(args.actual), 'rows_per_audit': 20789,
                      'cells_per_audit': 103945, 'server_sources': len(args.server_source)}))


if __name__ == '__main__':
    main()
