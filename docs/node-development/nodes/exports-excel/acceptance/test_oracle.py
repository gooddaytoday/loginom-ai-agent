"""Offline oracle qualification with an independent openpyxl writer, not Loginom."""
import copy
import datetime as dt
import json
from pathlib import Path
import tempfile
import unittest

from openpyxl import Workbook

from oracle import ROOT, check_fixture, verify


class OracleTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.path = Path(self.directory.name) / 'excel-cells.xlsx'
        self.expected = json.loads((ROOT / 'expected.json').read_text(encoding='utf-8'))
        self.book = Workbook()
        self.addCleanup(self.book.close)
        self.sheet = self.book.active
        self.sheet.title = 'Данные'
        self.sheet.append([c['label'] for c in self.expected['columns']])
        for values in self.expected['rows']:
            row = list(values)
            if row[3] is not None:
                row[3] = dt.datetime.fromisoformat(row[3])
            self.sheet.append(row)
        # Source value explicitly requests a string, not an Excel formula.
        self.sheet['B6'].data_type = 's'

    def result(self):
        self.book.save(self.path)
        return verify(self.path, self.expected)

    def test_complete_reference_and_source(self):
        check_fixture(self.expected)
        self.assertEqual(self.result()['cells'], 30)

    def test_rounded_number(self):
        self.sheet['C2'] = 1234.57
        with self.assertRaisesRegex(ValueError, 'CELL_VALUE_OR_TYPE_C2'):
            self.result()

    def test_boolean_is_not_integer(self):
        self.sheet['E2'] = 1
        with self.assertRaisesRegex(ValueError, 'CELL_VALUE_OR_TYPE_E2'):
            self.result()

    def test_null_is_not_empty_string(self):
        self.sheet['B4'] = ''
        with self.assertRaisesRegex(ValueError, 'CELL_VALUE_OR_TYPE_B4'):
            self.result()

    def test_empty_string_is_not_null(self):
        self.sheet['B5'] = None
        with self.assertRaisesRegex(ValueError, 'CELL_VALUE_OR_TYPE_B5'):
            self.result()

    def test_formula_is_not_text(self):
        self.sheet['B6'] = '=SUM(A1:A2)'
        with self.assertRaisesRegex(ValueError, 'FORMULA_B6'):
            self.result()

    def test_row_order(self):
        self.sheet['A2'], self.sheet['A3'] = 2, 1
        with self.assertRaisesRegex(ValueError, 'CELL_VALUE_OR_TYPE_A2'):
            self.result()

    def test_column_order(self):
        self.sheet['A1'], self.sheet['B1'] = 'Текст', 'Номер'
        with self.assertRaisesRegex(ValueError, 'CELL_VALUE_OR_TYPE_A1'):
            self.result()

    def test_extra_sheet(self):
        self.book.create_sheet('Extra')
        with self.assertRaisesRegex(ValueError, 'EXACTLY_ONE_SHEET'):
            self.result()

    def test_extra_cell(self):
        self.sheet['A7'] = 6
        with self.assertRaisesRegex(ValueError, 'EXTRA_CELLS'):
            self.result()

    def test_date_precision(self):
        self.sheet['D2'] = self.sheet['D2'].value + dt.timedelta(milliseconds=10)
        with self.assertRaisesRegex(ValueError, 'CELL_VALUE_OR_TYPE_D2'):
            self.result()

    def test_names_and_no_header_and_empty(self):
        original = copy.deepcopy(self.expected)
        for mode in ['names', 'none', 'empty']:
            self.expected = copy.deepcopy(original)
            header = 'labels' if mode == 'empty' else mode
            self.expected['header'] = header
            self.sheet.delete_rows(1, self.sheet.max_row)
            if header != 'none':
                self.sheet.append([c['name' if header == 'names' else 'label'] for c in self.expected['columns']])
            if mode == 'empty':
                self.expected['rows'] = []
                self.expected['row_count'] = 0
            for values in self.expected['rows']:
                row = list(values)
                row[3] = dt.datetime.fromisoformat(row[3]) if row[3] else None
                self.sheet.append(row)
            if self.expected['row_count']:
                self.sheet['B5' if header == 'none' else 'B6'].data_type = 's'
            self.assertEqual(self.result()['data_rows'], self.expected['row_count'])


if __name__ == '__main__':
    unittest.main()
