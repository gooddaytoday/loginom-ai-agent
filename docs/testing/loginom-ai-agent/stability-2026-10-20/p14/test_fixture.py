"""Run with dependencies from requirements.txt; no Loginom/network access."""
import csv
import json
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile
import unittest
from zipfile import ZipFile

import openpyxl

ROOT = Path(__file__).resolve().parent


class WorkbookRegression(unittest.TestCase):
    def test_standard_workbook_preserves_all_input_values(self):
        # The previous minimal OOXML passed XML readers but produced COL1 in
        # Loginom 7.4.2. Keep the complete standard-writer package here; live
        # CLI/cold runs remain the interoperability check.
        with ZipFile(ROOT / 'data/exhibits.xlsx') as book:
            self.assertIn('xl/styles.xml', book.namelist())
            self.assertIn('xl/sharedStrings.xml', book.namelist())
        book = openpyxl.load_workbook(ROOT / 'data/exhibits.xlsx', data_only=True)
        try:
            self.assertEqual(book.sheetnames, ['Exhibits'])
            sheet = book.worksheets[0]
            self.assertEqual((sheet.max_row, sheet.max_column), (16, 3))
            self.assertEqual(list(sheet.merged_cells), [])
            headers, *values = list(sheet.values)
            self.assertEqual(headers, ('Id', 'Gallery', 'Weight'))
            with (ROOT / 'data/exhibits.csv').open(newline='') as stream:
                source = list(csv.DictReader(stream))
            rows = [dict(Id=int(r['Id']), Gallery=r['Gallery'], Weight=float(r['Weight'])) for r in source]
            self.assertEqual([dict(zip(headers, row)) for row in values], rows)
            expected = json.loads((ROOT / 'expected-demo.json').read_text())
            self.assertEqual(rows, expected['outputs'][0]['rows'])
            actual = subprocess.check_output([sys.executable, str(ROOT / 'oracle.py')])
            self.assertEqual(actual, (ROOT / 'expected-demo.json').read_bytes())
        finally:
            book.close()

    def test_generator_is_reproducible_from_another_directory(self):
        with tempfile.TemporaryDirectory() as directory:
            case = Path(directory) / 'p14'
            case.mkdir()
            shutil.copyfile(ROOT / 'generate.py', case / 'generate.py')
            for _ in range(2):
                subprocess.run([sys.executable, str(case / 'generate.py')], cwd=directory, check=True)
                for name in ('exhibits.xlsx', 'exhibits.csv'):
                    self.assertEqual((case / 'data' / name).read_bytes(), (ROOT / 'data' / name).read_bytes())


if __name__ == '__main__':
    unittest.main()
